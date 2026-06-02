"""
Paper Formattor — 학생용 시험지 + 교수용 정답지 docx/pdf 생성.

LLM 호출 없음. python-docx 기반 결정론 렌더링.
PDF는 best-effort: docx2pdf 미설치/변환 실패 시 None 반환 + warning 누적.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from pydantic import Field

from common.schemas import (
    AnswerRubric,
    ExamBlueprint,
    Question,
    StrictBase,
)
from config.defaults import (
    EXAM_INSTRUCTIONS_KR,
    EXAM_METADATA,
    OUTPUT_DIR_TEMPLATE,
    PROJECT_ROOT,
)

logger = logging.getLogger(__name__)

# ────────────────────────────────────────────────────────────
# 상수
# ────────────────────────────────────────────────────────────
_TYPE_ORDER: dict[str, int] = {
    "short_answer": 0,
    "long_answer": 1,
    "case_analysis": 2,
    "MCQ_single": -1,
}

_SECTION_NAMES: dict[str, str] = {
    "short_answer": "단답형 (Short Answer)",
    "long_answer": "서술형 (Long Answer)",
    "case_analysis": "사례 분석 (Case Analysis)",
}

_ANSWER_SPACE_LINES: dict[str, int] = {
    "short_answer": 2,
    "long_answer": 6,
    "case_analysis": 8,
}


# ────────────────────────────────────────────────────────────
# 출력 결과 스키마
# ────────────────────────────────────────────────────────────
class FormatResult(StrictBase):
    session_id: str
    student_docx_path: str
    student_pdf_path: Optional[str] = None
    answer_key_docx_path: str
    answer_key_pdf_path: Optional[str] = None
    warnings: list[str] = Field(default_factory=list)


# ────────────────────────────────────────────────────────────
# 헬퍼
# ────────────────────────────────────────────────────────────
def _qtype_str(obj) -> str:
    """QuestionType enum 또는 str에서 문자열 값 반환."""
    return obj.value if hasattr(obj, "value") else str(obj)


def _bold_para(doc: Document, text: str, size_pt: Optional[int] = None) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    if size_pt:
        run.font.size = Pt(size_pt)


def _centered_bold(doc: Document, text: str, size_pt: int) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(size_pt)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ────────────────────────────────────────────────────────────
# PaperFormattor
# ────────────────────────────────────────────────────────────
class PaperFormattor:
    def __init__(self, base_dir: Optional[str] = None) -> None:
        self._base_dir = base_dir
        self._warnings: list[str] = []

    # ── 공개 인터페이스 ──────────────────────────────────────
    def format(
        self,
        blueprint: ExamBlueprint,
        questions: list[Question],
        rubrics: list[AnswerRubric],
        session_id: str,
    ) -> FormatResult:
        self._warnings = []

        # 매칭 및 정렬
        self._match_questions_to_slots(blueprint, questions)
        q_rubric_map = self._match_rubrics(questions, rubrics)
        ordered = self._order_questions_by_section(questions, blueprint)

        # 출력 디렉터리
        if self._base_dir:
            out_dir = Path(self._base_dir) / session_id
        else:
            out_dir = PROJECT_ROOT / OUTPUT_DIR_TEMPLATE.format(session_id=session_id)
        out_dir.mkdir(parents=True, exist_ok=True)

        # 학생용 docx
        student_docx = str(out_dir / "exam_student.docx")
        self._build_student_docx(blueprint, ordered, student_docx)
        student_pdf = self._try_convert_to_pdf(student_docx)

        # 교수용 docx
        answer_key_docx = str(out_dir / "exam_answer_key.docx")
        self._build_answer_key_docx(blueprint, ordered, q_rubric_map, answer_key_docx)
        answer_key_pdf = self._try_convert_to_pdf(answer_key_docx)

        return FormatResult(
            session_id=session_id,
            student_docx_path=student_docx,
            student_pdf_path=student_pdf,
            answer_key_docx_path=answer_key_docx,
            answer_key_pdf_path=answer_key_pdf,
            warnings=list(self._warnings),
        )

    # ── 내부: 매칭 ───────────────────────────────────────────
    def _match_questions_to_slots(
        self,
        blueprint: ExamBlueprint,
        questions: list[Question],
    ) -> dict[str, Question]:
        q_by_slot = {q.slot_id: q for q in questions}
        missing = [
            s.slot_id
            for s in blueprint.question_slots
            if s.slot_id not in q_by_slot
        ]
        if missing:
            raise ValueError(f"No question for slots: {missing}")
        return q_by_slot

    def _match_rubrics(
        self,
        questions: list[Question],
        rubrics: list[AnswerRubric],
    ) -> dict[str, AnswerRubric]:
        rubric_map = {r.question_id: r for r in rubrics}
        result: dict[str, AnswerRubric] = {}
        for q in questions:
            if q.question_id not in rubric_map:
                raise ValueError(f"No rubric for question {q.question_id}")
            result[q.question_id] = rubric_map[q.question_id]
        return result

    # ── 내부: 정렬 ───────────────────────────────────────────
    def _order_questions_by_section(
        self,
        questions: list[Question],
        blueprint: ExamBlueprint,
    ) -> list[Question]:
        slot_idx = {s.slot_id: i for i, s in enumerate(blueprint.question_slots)}

        def key(q: Question) -> tuple[int, int]:
            qtype = _qtype_str(q.question_type)
            return (_TYPE_ORDER.get(qtype, 9), slot_idx.get(q.slot_id, 999))

        return sorted(questions, key=key)

    # ── 내부: 헤더 빌드 (학생용/교수용 공통) ─────────────────
    def _add_header(self, doc: Document, answer_key: bool = False) -> None:
        title = (
            f"{EXAM_METADATA['course_name_kr']} ({EXAM_METADATA['course_name_en']}) "
            f"— {EXAM_METADATA['semester']} 학기 중간고사"
        )
        if answer_key:
            title += "  [ANSWER KEY (교수용)]"
        _centered_bold(doc, title, 14)

        # 메타 표
        meta_tbl = doc.add_table(rows=1, cols=3)
        try:
            meta_tbl.style = "Table Grid"
        except Exception:
            pass
        cells = meta_tbl.rows[0].cells
        cells[0].text = f"총점: {EXAM_METADATA['total_points']}점"
        cells[1].text = f"시간: {EXAM_METADATA['duration_minutes']}분"
        cells[2].text = f"{EXAM_METADATA['total_questions']}문항"
        for cell in cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()
        info = doc.add_paragraph("이름: ____________  학번: ____________")
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()
        _bold_para(doc, "지시사항", size_pt=12)
        for i, instr in enumerate(EXAM_INSTRUCTIONS_KR, 1):
            doc.add_paragraph(f"{i}. {instr}")

        doc.add_paragraph()
        sep = doc.add_paragraph("─" * 60)
        sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

    # ── 내부: 학생용 docx ─────────────────────────────────────
    def _build_student_docx(
        self,
        blueprint: ExamBlueprint,
        ordered_questions: list[Question],
        output_path: str,
    ) -> None:
        doc = Document()
        self._add_header(doc, answer_key=False)

        sections_sorted = sorted(
            blueprint.section_structure,
            key=lambda s: _TYPE_ORDER.get(_qtype_str(s.question_type), 9),
        )

        q_num = 1
        for sec_idx, section in enumerate(sections_sorted, 1):
            qtype = _qtype_str(section.question_type)
            sec_pts = section.num_questions * section.points_per_question
            sec_name = _SECTION_NAMES.get(qtype, qtype)

            # 2번째 섹션부터는 새 페이지에서 시작해 섹션 헤더를 첫 문제와 같이 배치
            if sec_idx > 1:
                doc.add_page_break()

            sec_para = doc.add_paragraph()
            run = sec_para.add_run(
                f"Section {sec_idx}. {sec_name} "
                f"({section.num_questions}문항 × {section.points_per_question}점 = {sec_pts}점)"
            )
            run.bold = True
            run.font.size = Pt(12)
            doc.add_paragraph()

            sec_questions = [
                q for q in ordered_questions if _qtype_str(q.question_type) == qtype
            ]
            space = _ANSWER_SPACE_LINES.get(qtype, 4)

            for q_idx, q in enumerate(sec_questions):
                # 같은 섹션 내 2번째 문제부터 페이지 나누기
                if q_idx > 0:
                    doc.add_page_break()

                q_para = doc.add_paragraph()
                q_run = q_para.add_run(f"문제 {q_num}. [{q.points}점]")
                q_run.bold = True

                text_para = doc.add_paragraph(q.prompt)
                if text_para.runs:
                    text_para.runs[0].font.size = Pt(11)

                for _ in range(space):
                    doc.add_paragraph("")

                doc.add_paragraph()
                q_num += 1

        doc.save(output_path)

    # ── 내부: 교수용 docx ─────────────────────────────────────
    def _build_answer_key_docx(
        self,
        blueprint: ExamBlueprint,
        ordered_questions: list[Question],
        q_rubric_map: dict[str, AnswerRubric],
        output_path: str,
    ) -> None:
        doc = Document()
        self._add_header(doc, answer_key=True)

        sections_sorted = sorted(
            blueprint.section_structure,
            key=lambda s: _TYPE_ORDER.get(_qtype_str(s.question_type), 9),
        )

        q_num = 1
        for sec_idx, section in enumerate(sections_sorted, 1):
            qtype = _qtype_str(section.question_type)
            sec_pts = section.num_questions * section.points_per_question
            sec_name = _SECTION_NAMES.get(qtype, qtype)

            # 2번째 섹션부터는 새 페이지에서 시작해 섹션 헤더를 첫 문제와 같이 배치
            if sec_idx > 1:
                doc.add_page_break()

            sec_para = doc.add_paragraph()
            run = sec_para.add_run(
                f"Section {sec_idx}. {sec_name} "
                f"({section.num_questions}문항 × {section.points_per_question}점 = {sec_pts}점)"
            )
            run.bold = True
            run.font.size = Pt(12)
            doc.add_paragraph()

            sec_questions = [
                q for q in ordered_questions if _qtype_str(q.question_type) == qtype
            ]

            for q_idx, q in enumerate(sec_questions):
                # 같은 섹션 내 2번째 문제부터 페이지 나누기
                if q_idx > 0:
                    doc.add_page_break()

                # 문제 번호 + 배점
                q_para = doc.add_paragraph()
                q_run = q_para.add_run(f"문제 {q_num}. [{q.points}점]")
                q_run.bold = True

                # 문제 텍스트
                text_para = doc.add_paragraph(q.prompt)
                if text_para.runs:
                    text_para.runs[0].font.size = Pt(11)

                # 모범 답안
                _bold_para(doc, "▶ 모범답안:")
                doc.add_paragraph(q.reference_answer)

                # 채점 기준
                _bold_para(doc, "▶ 채점 기준:")
                rubric = q_rubric_map[q.question_id]
                self._add_rubric_table(doc, rubric)

                doc.add_paragraph()
                q_num += 1

        doc.save(output_path)

    def _add_rubric_table(self, doc: Document, rubric: AnswerRubric) -> None:
        n_criteria = len(rubric.criteria)
        table = doc.add_table(rows=1 + n_criteria + 1, cols=5)
        try:
            table.style = "Table Grid"
        except Exception:
            pass

        # 열 폭 설정 (A4 가용 폭 ≈ 15.9cm)
        # 기준 ID·배점은 좁게, 핵심 요소·부분점수 기준은 넓게
        col_widths = [Cm(1.2), Cm(3.2), Cm(1.2), Cm(5.0), Cm(5.3)]
        for col_idx, width in enumerate(col_widths):
            for cell in table.column_cells(col_idx):
                cell.width = width

        # 헤더 행
        headers = ["기준 ID", "평가 항목", "배점", "핵심 요소", "부분 점수 기준"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            if cell.paragraphs and cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].bold = True

        # 기준 행
        for i, crit in enumerate(rubric.criteria):
            row = table.rows[i + 1]
            row.cells[0].text = crit.criterion_id
            row.cells[1].text = crit.description
            row.cells[2].text = str(crit.points)
            row.cells[3].text = "\n".join(crit.key_points) if crit.key_points else ""
            row.cells[4].text = crit.partial_credit_guide

        # 합계 행
        total_row = table.rows[-1]
        total_row.cells[0].text = "합계"
        total_row.cells[2].text = str(rubric.total_points)

    # ── 내부: PDF 변환 ────────────────────────────────────────
    def _try_convert_to_pdf(self, docx_path: str) -> Optional[str]:
        try:
            from docx2pdf import convert  # type: ignore
        except ImportError:
            self._add_warning(
                f"docx2pdf not installed; PDF skipped for {docx_path}"
            )
            return None
        try:
            pdf_path = docx_path.replace(".docx", ".pdf")
            convert(docx_path, pdf_path)
            return pdf_path
        except Exception as exc:
            self._add_warning(f"PDF conversion failed for {docx_path}: {exc}")
            return None

    def _add_warning(self, msg: str) -> None:
        self._warnings.append(msg)
        logger.warning(msg)
